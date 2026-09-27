"""SP3 — amostra de divergências para a tipologia, com categoria sugerida por heurística.

Seleção (720 itens ofensivos, sem as 16 cópias exatas, texto sem prefixo):
- Direção A (métrica alta, humano baixo): plausibilidade <= 1/3, os 50 maiores ROUGE-L.
- Direção B (métrica baixa, humano alto): plausibilidade >= 2/3, os 50 menores ROUGE-L
  (desempate: maior plausibilidade).

As categorias sugeridas são só um ponto de partida para a anotação manual do Caio;
a concordância entre sugestão e decisão final é calculada na aba "Resumo".

Entrada : results/metricas_por_item.csv, results/metricas_multiref.csv
Saída   : results/sp3_amostra.csv, results/sp3_anotacao.xlsx

Uso:  python src/amostra_sp3.py
"""
from __future__ import annotations

import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from config import METRICAS_CSV, PREFIXO, RESULTADOS

N_POR_DIRECAO = 50

CATEGORIAS = [
    ("A1", "A", "Grupo certo, estereótipo errado",
     "Acerta o grupo-alvo, mas atribui a ele um estereótipo diferente do que o post implica.",
     "ref: black folks are violent · gerada: black folks are stupid"),
    ("A2", "A", "Sobreposição só de molde",
     "As palavras em comum são estrutura ou molde (\"is a personal attack\", \"can't\"), não conteúdo.",
     "ref: this post is a personal attack · gerada: incel redemption story is a personal attack"),
    ("A3", "A", "Sentido invertido ou positivo",
     "Mesmo grupo, mas a gerada diz algo positivo ou oposto ao estereótipo.",
     "ref: women are cowardly · gerada: women are smarter"),
    ("A4", "A", "Referência também implausível",
     "Os humanos também deram nota baixa à referência: o problema é o item, não a métrica.",
     "plausibilidade da referência <= 1/3"),
    ("A9", "A", "Outro (A)", "Nenhuma das anteriores. Descreva na observação.", ""),
    ("B1", "B", "Paráfrase",
     "Mesma ideia com outras palavras (inclui nomes diferentes para o mesmo grupo).",
     "ref: islam is violent · gerada: muslims are bad"),
    ("B2", "B", "Explicação genérica",
     "A referência (ou a gerada) não nomeia grupo nem estereótipo: \"personal attack\", \"trivializes harm\".",
     "ref: this post is a personal attack · gerada: women are objects"),
    ("B3", "B", "Implicação alternativa coberta",
     "Outra leitura válida do post, que aparece em outra referência do SBIC (ROUGE-L múltiplo alto).",
     "ref: nobody cares if they drown · gerada: mexicans can't swim (ROUGE-L múlt. = 1)"),
    ("B4", "B", "Implicação alternativa não coberta",
     "Outra leitura válida do post, que nenhuma referência do SBIC captura.",
     "ref: fat folks don't get love · gerada: women are more attractive than men"),
    ("B9", "B", "Outro (B)", "Nenhuma das anteriores. Descreva na observação.", ""),
]

GENERICO = re.compile(
    r"personal attack|trivializes|no stereotype|makes light|not offensive|jokes? about|"
    r"does not imply|punchline",
    re.I,
)
POSITIVO = {"smarter", "happy", "good", "better", "nice", "exploited", "victims", "oppressed",
            "kind", "beautiful", "strong", "successful", "superior"}
VERBOS = r"\b(are|is|can|can't|cannot|should|shouldn't|don't|do|does|doesn't|have|has|aren't|isn't|were|was|will|won't|like|deserve|need|must|get|make)\b"
VAZIAS = {"folks", "people", "person", "persons", "the", "all", "a", "an", "of", "with", "who", "that", "men", "guys"}
SINONIMOS = {
    "jewish": "jew", "jews": "jew", "blacks": "black", "african": "black", "americans": "american",
    "muslims": "muslim", "islam": "muslim", "islamic": "muslim", "arabs": "arab", "catholics": "catholic",
    "priests": "catholic", "christians": "christian", "lesbians": "lesbian", "gays": "gay",
    "homosexuals": "gay", "homosexual": "gay", "mexicans": "mexican", "latinos": "mexican",
    "hispanic": "mexican", "females": "women", "woman": "women", "girls": "women", "whites": "white",
    "asians": "asian", "chinese": "asian", "minors": "children", "kids": "children", "child": "children",
    "immigrants": "immigrant", "refugees": "immigrant", "trans": "transgender", "transgenders": "transgender",
}


def sem_prefixo(t: str) -> str:
    return re.sub(PREFIXO, "", str(t).strip().lower()).rstrip(". ")


def grupo(t: str) -> set[str]:
    """Palavras antes do primeiro verbo, normalizadas (heurística do 'grupo-alvo')."""
    partes = re.split(VERBOS, sem_prefixo(t), maxsplit=1)
    palavras = re.findall(r"[a-z]+", partes[0])
    return {SINONIMOS.get(w, w.rstrip("s") if len(w) > 4 else w) for w in palavras} - VAZIAS


def resto(t: str) -> set[str]:
    partes = re.split(VERBOS, sem_prefixo(t), maxsplit=1)
    return set(re.findall(r"[a-z]+", partes[-1])) if len(partes) > 1 else set()


def sugerir(r: pd.Series, direcao: str) -> tuple[str, str, str]:
    g_ref, g_ger = grupo(r.referencia), grupo(r.gerada)
    mesmo_grupo = bool(g_ref & g_ger)
    generica = bool(GENERICO.search(str(r.referencia)) or GENERICO.search(str(r.gerada)))
    if direcao == "A":
        if r.plausibilidade_ref <= 1 / 3 + 1e-9:
            extra = ""
            if mesmo_grupo:
                extra = "; também tem padrão A3" if resto(r.gerada) & POSITIVO else "; também tem padrão A1"
            return "A4", f"plausibilidade da referência = {r.plausibilidade_ref:.2f}{extra}", "alta"
        if generica:
            return "A2", "referência ou gerada usa molde genérico", "média"
        if mesmo_grupo and resto(r.gerada) & POSITIVO:
            return "A3", f"mesmo grupo ({', '.join(sorted(g_ref & g_ger))}); termo positivo na gerada", "média"
        if mesmo_grupo:
            return "A1", f"mesmo grupo ({', '.join(sorted(g_ref & g_ger))}), estereótipo diferente", "alta"
        return "A9", "sem padrão reconhecido", "baixa"
    if GENERICO.search(str(r.referencia)):
        return "B2", "referência genérica (não nomeia grupo/estereótipo)", "alta"
    if GENERICO.search(str(r.gerada)):
        return "B2", "gerada genérica, mas julgada plausível", "média"
    if mesmo_grupo:
        return "B1", f"mesmo grupo ({', '.join(sorted(g_ref & g_ger))}), palavras diferentes", "média"
    if r.rougeL_multiref_semprefixo >= 0.5:
        return "B3", f"ROUGE-L com múltiplas refs = {r.rougeL_multiref_semprefixo:.2f}", "média"
    return "B4", "grupos diferentes e nenhuma referência do SBIC cobre", "baixa"


def selecionar() -> pd.DataFrame:
    df = pd.read_csv(METRICAS_CSV).merge(
        pd.read_csv(RESULTADOS / "metricas_multiref.csv"), on=["modelo", "ques_id"], how="left"
    )
    of = df[df.ofensivo & ~df.copia_exata]
    a = of[of.plausibilidade <= 1 / 3 + 1e-9].sort_values(
        ["rougeL_semprefixo", "plausibilidade"], ascending=[False, True]).head(N_POR_DIRECAO)
    b = of[of.plausibilidade >= 2 / 3 - 1e-9].sort_values(
        ["rougeL_semprefixo", "plausibilidade"], ascending=[True, False]).head(N_POR_DIRECAO)
    linhas = []
    for direcao, bloco in [("A", a), ("B", b)]:
        for i, (_, r) in enumerate(bloco.iterrows(), 1):
            cat, motivo, conf = sugerir(r, direcao)
            linhas.append({
                "id": f"{direcao}-{i:02d}", "direcao": direcao, "modelo": r.modelo, "ques_id": r.ques_id,
                "post": r.post, "referencia": r.referencia, "gerada": r.gerada,
                "plausibilidade": round(r.plausibilidade, 3), "notas": r.notas,
                "plausibilidade_ref": round(r.plausibilidade_ref, 3),
                "rougeL_unica": round(r.rougeL_semprefixo, 3),
                "rougeL_multipla": round(r.rougeL_multiref_semprefixo, 3), "n_refs": int(r.n_refs),
                "categoria_sugerida": cat, "motivo": motivo, "confianca": conf,
            })
    return pd.DataFrame(linhas)


# ----------------------------------------------------------------------------- planilha
FONTE = "Arial"
AZUL_ESC = "1F3864"
AMARELO = PatternFill("solid", fgColor="FFF2CC")
CINZA = PatternFill("solid", fgColor="F2F2F2")
CAB = PatternFill("solid", fgColor=AZUL_ESC)
FINA = Side(style="thin", color="BFBFBF")
BORDA = Border(left=FINA, right=FINA, top=FINA, bottom=FINA)


def f(bold=False, cor="000000", tam=10, italico=False):
    return Font(name=FONTE, bold=bold, color=cor, size=tam, italic=italico)


def aba_instrucoes(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Instruções"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 70
    ws.column_dimensions["E"].width = 62
    linha = 1
    ws.cell(linha, 1, "SP3 — Tipologia de divergência entre métricas e julgamento humano").font = f(True, AZUL_ESC, 14)
    linha += 1
    ws.cell(linha, 1, "PIBIC 2026–2027 · amostra de 100 itens do FEB/SBIC (50 por direção) · gerada por src/amostra_sp3.py").font = f(italico=True, cor="595959")
    linha += 2
    textos = [
        ("Aviso", "Os posts são do SBIC e contêm linguagem ofensiva e discurso de ódio. É o objeto de estudo; leia com isso em mente."),
        ("O que fazer", "Na aba \"Anotação\", preencha só as colunas amarelas: CATEGORIA FINAL (lista suspensa) e Observação. A categoria sugerida vem de uma heurística automática: se concordar, escolha o mesmo código; se não, escolha outro."),
        ("Direção A", "Métrica alta, humano baixo (plausibilidade ≤ 1/3, maiores ROUGE-L): a métrica \"aprova\" o que os humanos reprovam."),
        ("Direção B", "Métrica baixa, humano alto (plausibilidade ≥ 2/3, menores ROUGE-L): a métrica \"reprova\" o que os humanos aprovam."),
        ("Colunas", "Plausibilidade = média de 3 anotadores (yes = 1, w_yes = 2/3, w_no = 1/3, no = 0). ROUGE-L sem o prefixo \"this post implies that\". \"Múltipla\" = contra todas as referências do SBIC."),
        ("Dica", "Se ficar em dúvida entre duas categorias, escolha a principal e cite a outra na Observação. Use A9/B9 só quando nada servir."),
    ]
    for rot, txt in textos:
        ws.cell(linha, 1, rot).font = f(True)
        c = ws.cell(linha, 2, txt)
        c.font = f()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=5)
        ws.row_dimensions[linha].height = 30
        linha += 1
    linha += 1
    ws.cell(linha, 1, "Livro de códigos").font = f(True, AZUL_ESC, 12)
    linha += 1
    for j, h in enumerate(["Código", "Direção", "Categoria", "Definição", "Exemplo"], 1):
        c = ws.cell(linha, j, h)
        c.font = f(True, "FFFFFF")
        c.fill = CAB
        c.border = BORDA
    linha += 1
    inicio = linha
    for cod, dirc, nome, defin, ex in CATEGORIAS:
        for j, v in enumerate([cod, dirc, nome, defin, ex], 1):
            c = ws.cell(linha, j, v)
            c.font = f(bold=(j == 1))
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDA
        ws.row_dimensions[linha].height = 30
        linha += 1
    wb.defined_names["codigos"] = DefinedName(
        "codigos", attr_text=f"'Instruções'!$A${inicio}:$A${linha - 1}")
    linha += 1
    ws.cell(linha, 1, "Exemplo de linha preenchida").font = f(True, AZUL_ESC, 12)
    linha += 1
    ex = [("Referência", "this post implies that black folks are violent"),
          ("Gerada", "this post implies that black folks are stupid"),
          ("Humano / ROUGE-L", "0,00 / 0,75"),
          ("Sugerida", "A1 (mesmo grupo: black, estereótipo diferente)"),
          ("CATEGORIA FINAL", "A1"),
          ("Observação", "Ambos são estereótipos reais sobre o grupo; o post implica violência, não inteligência.")]
    for rot, v in ex:
        ws.cell(linha, 1, rot).font = f(True)
        c = ws.cell(linha, 2, v)
        c.font = f()
        if rot in ("CATEGORIA FINAL", "Observação"):
            c.fill = AMARELO
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=4)
        linha += 1


COLUNAS = [  # (título, chave, largura, formato)
    ("ID", "id", 7, None), ("Dir.", "direcao", 5, None), ("Modelo", "modelo", 9, None),
    ("Post", "post", 48, None), ("Referência (FEB)", "referencia", 32, None),
    ("Gerada", "gerada", 32, None), ("Plaus. humana", "plausibilidade", 9, "0.00"),
    ("Notas (3 anot.)", "notas", 14, None), ("Plaus. da ref.", "plausibilidade_ref", 9, "0.00"),
    ("ROUGE-L única", "rougeL_unica", 9, "0.00"), ("ROUGE-L múltipla", "rougeL_multipla", 9, "0.00"),
    ("Nº refs", "n_refs", 6, "0"), ("Sugerida", "categoria_sugerida", 9, None),
    ("Motivo da sugestão", "motivo", 30, None), ("Confiança", "confianca", 9, None),
    ("CATEGORIA FINAL", None, 11, None), ("Observação", None, 36, None),
]


def aba_anotacao(wb: Workbook, df: pd.DataFrame) -> None:
    ws = wb.create_sheet("Anotação")
    for j, (tit, _, larg, _) in enumerate(COLUNAS, 1):
        c = ws.cell(1, j, tit)
        c.font = f(True, "FFFFFF")
        c.fill = CAB if j < 16 else PatternFill("solid", fgColor="BF8F00")
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        c.border = BORDA
        ws.column_dimensions[get_column_letter(j)].width = larg
    ws.row_dimensions[1].height = 32
    ws.cell(1, 16).comment = Comment("Escolha um código da lista (A1–A4, A9 para direção A; B1–B4, B9 para B).", "PIBIC")
    for i, r in enumerate(df.itertuples(index=False), 2):
        faixa = CINZA if r.direcao == "B" else None
        for j, (_, chave, _, fmt) in enumerate(COLUNAS, 1):
            v = getattr(r, chave) if chave else None
            c = ws.cell(i, j, v)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDA
            if fmt:
                c.number_format = fmt
            if j >= 16:
                c.fill = AMARELO
            elif faixa:
                c.fill = faixa
        ws.row_dimensions[i].height = 60
    ult = len(df) + 1
    dv = DataValidation(type="list", formula1="=codigos", allow_blank=True,
                        error="Use um código do livro de códigos (aba Instruções).", errorTitle="Código inválido")
    ws.add_data_validation(dv)
    dv.add(f"P2:P{ult}")
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:Q{ult}"


def aba_resumo(wb: Workbook, n: int) -> None:
    ws = wb.create_sheet("Resumo")
    ws.sheet_view.showGridLines = False
    ult = n + 1
    rng = lambda col: f"'Anotação'!${col}$2:${col}${ult}"
    ws["A1"] = "Resumo da anotação (atualiza sozinho)"
    ws["A1"].font = f(True, AZUL_ESC, 14)
    ws["A3"], ws["B3"] = "Itens anotados", f"=COUNTA({rng('P')})"
    ws["A4"], ws["B4"] = "Progresso", f"=B3/{n}"
    ws["B4"].number_format = "0%"
    ws["A5"], ws["B5"] = "Concordância com a sugestão", (
        f"=IFERROR(SUMPRODUCT(({rng('M')}={rng('P')})*({rng('P')}<>\"\"))/B3,\"-\")")
    ws["B5"].number_format = "0%"
    for a in ("A3", "A4", "A5"):
        ws[a].font = f(True)
    for b in ("B3", "B4", "B5"):
        ws[b].font = f()
    ws["C5"] = "Só conta itens já anotados."
    ws["C5"].font = f(italico=True, cor="595959")
    cab = ["Código", "Categoria", "Sugeridas", "Final", "% da direção (final)"]
    for j, h in enumerate(cab, 1):
        c = ws.cell(7, j, h)
        c.font = f(True, "FFFFFF")
        c.fill = CAB
        c.border = BORDA
    for i, (cod, dirc, nome, *_ ) in enumerate(CATEGORIAS, 8):
        ws.cell(i, 1, cod)
        ws.cell(i, 2, nome)
        ws.cell(i, 3, f"=COUNTIF({rng('M')},A{i})")
        ws.cell(i, 4, f"=COUNTIF({rng('P')},A{i})")
        ws.cell(i, 5, f"=IFERROR(D{i}/COUNTIFS({rng('B')},\"{dirc}\",{rng('P')},\"<>\"),\"-\")")
        ws.cell(i, 5).number_format = "0%"
        for j in range(1, 6):
            ws.cell(i, j).font = f()
            ws.cell(i, j).border = BORDA
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 36
    for col in "CDE":
        ws.column_dimensions[col].width = 16


def main() -> None:
    df = selecionar()
    df.to_csv(RESULTADOS / "sp3_amostra.csv", index=False)
    wb = Workbook()
    aba_instrucoes(wb)
    aba_anotacao(wb, df)
    aba_resumo(wb, len(df))
    wb.save(RESULTADOS / "sp3_anotacao.xlsx")
    print(df.groupby(["direcao", "categoria_sugerida"]).size().to_string())
    print(df.groupby("confianca").size().to_string())
    print(f"-> {RESULTADOS / 'sp3_anotacao.xlsx'}")


if __name__ == "__main__":
    main()
