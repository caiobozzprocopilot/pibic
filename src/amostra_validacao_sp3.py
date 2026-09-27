"""SP3 — planilha para a anotação humana independente (validação, pedido do Orientador).

Sorteia 40 dos 100 casos (20 da direção A, 20 da B; semente fixa) e gera uma planilha
SEM a sugestão da heurística (M–O) e SEM a anotação do Claude (P–R). O livro de códigos
é o original (A1–A4, A9, B1–B4, B9); A5 "leitura literal" NÃO está disponível (P11).

Saída: results/sp3_validacao_humana.xlsx (anotar na coluna "Categoria (humano)")
Depois: python src/analise_sp3.py -> κ Claude × humano nos 40 itens
Uso:    python src/amostra_validacao_sp3.py
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from amostra_sp3 import AMARELO, BORDA, CAB, CINZA, aba_instrucoes, f
from config import RESULTADOS, SEMENTE

N_POR_DIRECAO = 20
COLS = [("ID", "id", 7), ("Dir.", "direcao", 5), ("Modelo", "modelo", 9), ("Post", "post", 48),
        ("Referência (FEB)", "referencia", 32), ("Gerada", "gerada", 32),
        ("Plaus. humana", "plausibilidade", 9), ("Notas (3 anot.)", "notas", 14),
        ("Plaus. da ref.", "plausibilidade_ref", 9), ("ROUGE-L única", "rougeL_unica", 9),
        ("ROUGE-L múltipla", "rougeL_multipla", 9), ("Nº refs", "n_refs", 6),
        ("Categoria (humano)", None, 12), ("Observação", None, 40)]


def main() -> None:
    amostra = pd.read_csv(RESULTADOS / "sp3_amostra.csv")
    sel = pd.concat([g.sample(N_POR_DIRECAO, random_state=SEMENTE) for _, g in amostra.groupby("direcao")])
    sel = sel.sample(frac=1, random_state=SEMENTE)  # intercala A e B em ordem aleatória

    wb = Workbook()
    aba_instrucoes(wb)
    wi = wb["Instruções"]
    wi.insert_rows(3)
    c = wi.cell(3, 1, "VALIDAÇÃO INDEPENDENTE: anote sem abrir sp3_anotacao_cega_final.xlsx nem "
                      "sp3_anotacao_claude.tsv até terminar. Preencha só 'Categoria (humano)' e, se quiser, "
                      "'Observação'. A5 não existe neste livro de códigos: se achar um padrão novo, use A9/B9 "
                      "e descreva na Observação.")
    c.font = f(True, "C00000")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    wi.merge_cells(start_row=3, start_column=1, end_row=3, end_column=5)
    wi.row_dimensions[3].height = 45
    # insert_rows não atualiza nomes definidos: reaponta "codigos" para as linhas dos códigos
    linhas_cod = [i for i in range(1, wi.max_row + 1)
                  if isinstance(wi.cell(i, 1).value, str) and len(wi.cell(i, 1).value) == 2
                  and wi.cell(i, 1).value[0] in "AB" and wi.cell(i, 1).value[1].isdigit()]
    from openpyxl.workbook.defined_name import DefinedName
    del wb.defined_names["codigos"]
    wb.defined_names["codigos"] = DefinedName(
        "codigos", attr_text=f"'Instruções'!$A${min(linhas_cod)}:$A${max(linhas_cod)}")

    ws = wb.create_sheet("Anotação")
    for j, (tit, _, larg) in enumerate(COLS, 1):
        c = ws.cell(1, j, tit)
        c.font = f(True, "FFFFFF")
        c.fill = CAB if j < 13 else PatternFill("solid", fgColor="BF8F00")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BORDA
        ws.column_dimensions[get_column_letter(j)].width = larg
    for i, r in enumerate(sel.itertuples(index=False), 2):
        for j, (_, chave, _) in enumerate(COLS, 1):
            c = ws.cell(i, j, getattr(r, chave) if chave else None)
            c.font = f()
            c.alignment = Alignment(wrap_text=True, vertical="top")
            c.border = BORDA
            if j >= 13:
                c.fill = AMARELO
            elif r.direcao == "B":
                c.fill = CINZA
            if chave in ("plausibilidade", "plausibilidade_ref", "rougeL_unica", "rougeL_multipla"):
                c.number_format = "0.00"
        ws.row_dimensions[i].height = 60
    ult = len(sel) + 1
    dv = DataValidation(type="list", formula1="=codigos", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"M2:M{ult}")
    ws.freeze_panes = "D2"

    wr = wb.create_sheet("Resumo")
    wr["A1"], wr["B1"] = "Itens anotados", f"=COUNTA('Anotação'!M2:M{ult})"
    wr["A2"], wr["B2"] = "Total", len(sel)
    for a in ("A1", "A2"):
        wr[a].font = f(True)
    wr.column_dimensions["A"].width = 20

    saida = RESULTADOS / "sp3_validacao_humana.xlsx"
    wb.save(saida)
    print(sel.groupby("direcao").size().to_string(), f"\n-> {saida}")


if __name__ == "__main__":
    main()
