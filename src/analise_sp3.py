"""SP3 — consolida a anotação da tipologia e mede concordância.

Lê results/sp3_anotacao_cega.xlsx (aba "Anotação"):
- M  "Sugerida"            -> heurística de src/amostra_sp3.py
- P  "CATEGORIA FINAL"     -> anotação do Caio (humana; a que vale para o relatório)
- R  "Categoria (Claude)"  -> segundo anotador (Claude), com justificativa em S
- Q  "Observação" = "cego" -> item anotado pelo Caio sem ver M–O

Calcula distribuição por categoria e concordância (% e kappa de Cohen) entre
heurística, Claude e Caio, separando itens às cegas dos demais (P9).

Saída: results/sp3_resumo.csv, results/sp3_concordancia.csv
Uso:   python src/analise_sp3.py
"""
from __future__ import annotations

import pandas as pd
from openpyxl import load_workbook

from config import RESULTADOS

PLANILHA = RESULTADOS / "sp3_anotacao_cega.xlsx"


def kappa(a: pd.Series, b: pd.Series) -> float:
    """Kappa de Cohen para dois rótulos nominais (sem dependência do sklearn)."""
    n = len(a)
    if n == 0:
        return float("nan")
    po = (a.values == b.values).mean()
    cats = set(a) | set(b)
    pe = sum((a == c).mean() * (b == c).mean() for c in cats)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def ler() -> pd.DataFrame:
    ws = load_workbook(PLANILHA, data_only=True)["Anotação"]
    linhas = [[c.value for c in r] for r in ws.iter_rows()]
    df = pd.DataFrame(linhas[1:], columns=linhas[0])
    df = df[df["ID"].notna()].copy()
    df["direcao"] = df["ID"].str[0]
    df["cego"] = df["Observação"].astype(str).str.lower().str.startswith("cego")
    df = df.rename(columns={"Sugerida": "heuristica", "CATEGORIA FINAL": "caio",
                            "Categoria (Claude)": "claude"})
    return df


def main() -> None:
    df = ler()
    anotadores = [c for c in ["heuristica", "claude", "caio"] if df[c].notna().any()]

    # distribuição
    dist = pd.concat(
        {a: df.groupby(["direcao", a]).size() for a in anotadores}, axis=1
    ).fillna(0).astype(int)
    dist.index.names = ["direcao", "categoria"]
    dist.to_csv(RESULTADOS / "sp3_resumo.csv")

    # concordância par a par, total e por subconjunto (cego / não cego)
    linhas = []
    for i, a in enumerate(anotadores):
        for b in anotadores[i + 1:]:
            for nome, sub in [("todos", df), ("cego", df[df.cego]), ("nao_cego", df[~df.cego])]:
                s = sub[sub[a].notna() & sub[b].notna()]
                linhas.append({"par": f"{a} x {b}", "subconjunto": nome, "n": len(s),
                               "concordancia": (s[a] == s[b]).mean() if len(s) else float("nan"),
                               "kappa": kappa(s[a], s[b])})
    conc = pd.DataFrame(linhas)
    conc.to_csv(RESULTADOS / "sp3_concordancia.csv", index=False)

    pd.set_option("display.float_format", "{:.2f}".format)
    print("Anotadores com dados:", ", ".join(anotadores))
    print(dist.to_string())
    print()
    print(conc.to_string(index=False))
    if "caio" not in anotadores:
        print("\nAviso: coluna CATEGORIA FINAL (Caio) ainda vazia.")


if __name__ == "__main__":
    main()
